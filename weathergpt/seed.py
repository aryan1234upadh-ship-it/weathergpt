"""Load your team's soil and crop sheets into the database.

Reads data/agri_data.xlsx (sheets SOIL and CROPS). If that file does not exist, it falls back
to data/soil_health.csv and data/region_crops.csv. Empty cells stay empty, nothing is guessed.
Run:  python seed.py
"""
import csv
import sys
from pathlib import Path

from app import create_app
from app.models import db, RegionCrop, SoilHealth

DATA = Path("data")
XLSX = DATA / "agri_data.xlsx"

SOIL_LIMITS = {"ph": (3, 10), "n": (0, 2000), "p": (0, 300), "k": (0, 3000), "moisture": (0, 100)}
CROP_LIMITS = {"t_min": (-10, 60), "t_max": (-10, 60), "rain_min": (0, 5000), "ph_min": (3, 10), "ph_max": (3, 10)}

problems, no_source = [], []
MAX_SHOWN = 30


def show(lines):
    for p in lines[:MAX_SHOWN]:
        print("  -", p)
    if len(lines) > MAX_SHOWN:
        print(f"  ... and {len(lines) - MAX_SHOWN} more")


def looks_like_garbage(d, number_keys, text_keys):
    """True when a row looks like text or code pasted into the wrong place."""
    bad = 0
    for k in number_keys:
        raw = d.get(k)
        if raw is None or str(raw).strip() == "":
            continue
        try:
            float(str(raw).replace(",", "").strip())
        except ValueError:
            bad += 1
    odd = any(ch in text(d, k) for k in text_keys for ch in '"[]{};=<>')
    return bad >= 2 or odd


def read_rows(sheet, csv_name):
    """Yield (row_number, dict) for every row that is not empty."""
    if XLSX.exists():
        try:
            from openpyxl import load_workbook
        except ImportError:
            sys.exit("Please run:  pip install openpyxl   and then run seed.py again.")
        wb = load_workbook(XLSX, data_only=True)
        if sheet not in wb.sheetnames:
            sys.exit(f"Sheet '{sheet}' was not found in {XLSX}")
        ws = wb[sheet]
        header = [str(c.value).strip().lower() if c.value is not None else "" for c in ws[1]]
        for i, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            d = {h: v for h, v in zip(header, row) if h}
            if any(v not in (None, "") for v in d.values()):
                yield i, d
    else:
        path = DATA / csv_name
        if not path.exists():
            return
        with open(path, newline="", encoding="utf-8-sig") as f:
            for i, d in enumerate(csv.DictReader(f), start=2):
                d = {(k or "").strip().lower(): v for k, v in d.items()}
                if any(v not in (None, "") for v in d.values()):
                    yield i, d


def text(d, key):
    v = d.get(key)
    return "" if v is None else str(v).strip()


def number(sheet, row, d, key, limits, who=""):
    raw = d.get(key)
    if raw is None or str(raw).strip() == "":
        return None
    try:
        value = float(str(raw).replace(",", "").strip())
    except ValueError:
        problems.append(f"{sheet} row {row}{who}: '{raw}' in column {key} is not a number, left empty")
        return None
    lo, hi = limits[key]
    if not lo <= value <= hi:
        problems.append(f"{sheet} row {row}{who}: {key} = {value} looks wrong (expected {lo} to {hi}), left empty")
        return None
    return value


soil, crops = [], []
if XLSX.exists():
    print(f"Reading {XLSX}")
else:
    print(f"NOTICE: {XLSX} was NOT found, so the old CSV files in {DATA} are being read instead.")
    print("        Put your filled Excel file there with exactly that name to use it.")

for row, d in read_rows("SOIL", "soil_health.csv"):
    if looks_like_garbage(d, SOIL_LIMITS, ["district", "state"]):
        problems.append(f"SOIL row {row}: looks like text pasted into the wrong place, row skipped")
        continue
    district = text(d, "district").title()
    if not district:
        problems.append(f"SOIL row {row}: no district, row skipped")
        continue
    vals = {k: number("SOIL", row, d, k, SOIL_LIMITS, f" ({district})") for k in SOIL_LIMITS}
    if all(v is None for v in vals.values()):
        problems.append(f"SOIL row {row} ({district}): no numbers, row skipped")
        continue
    source = text(d, "source_name") or text(d, "source")
    if not source:
        no_source.append(f"SOIL row {row} ({district})")
    soil.append(SoilHealth(district=district, state=text(d, "state") or None,
                           source=(source or "no source given")[:100], **vals))

for row, d in read_rows("CROPS", "region_crops.csv"):
    if looks_like_garbage(d, CROP_LIMITS, ["crop", "state", "region"]):
        problems.append(f"CROPS row {row}: looks like text pasted into the wrong place, row skipped")
        continue
    name, state = text(d, "crop"), text(d, "state")
    if not name or not state:
        problems.append(f"CROPS row {row}: crop and state are required, row skipped")
        continue
    vals = {k: number("CROPS", row, d, k, CROP_LIMITS, f" ({name})") for k in CROP_LIMITS}
    if all(v is None for v in vals.values()):
        problems.append(f"CROPS row {row} ({name}): no numbers, row skipped")
        continue
    if not text(d, "source_name"):
        no_source.append(f"CROPS row {row} ({name}, {state})")
    crops.append(RegionCrop(region=text(d, "region"), state=state, crop=name,
                            irrigation=text(d, "irrigation") or None, tips=text(d, "tips") or None, **vals))

app = create_app()
with app.app_context():
    if soil:
        SoilHealth.query.delete()
        for s in soil:
            db.session.add(s)
    else:
        print("No usable SOIL rows found. Keeping the soil data already in the database.")
    if crops:
        RegionCrop.query.delete()
        for c in crops:
            db.session.add(c)
    else:
        print("No usable CROPS rows found. Keeping the crop data already in the database.")
    db.session.commit()

print(f"\nLoaded {len(soil)} soil rows and {len(crops)} crop rows.")
if problems:
    print("\nRows or cells that could not be used (fix them in the sheet and run again):")
    show(problems)
if no_source:
    print("\nRows WITHOUT a source (add source_name and source_link):")
    show(no_source)
if (soil or crops) and not problems and not no_source:
    print("All rows have a source and every number is in a sensible range.")